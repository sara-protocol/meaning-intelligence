# -*- coding: utf-8 -*-
"""生成《意义智能 · 元智能理论与分层应用》中英双语公开发布白皮书 DOCX。

排版严格对齐《意义智能》V7.0 正典（从 KDP 中文版实测提取）：
  A4，页边距 上下2.54cm / 左右3.17cm
  标题 黑体 #1B3A5C；正文 宋体 12pt #222222 首行缩进 1.5 倍行距
  英文正文 Georgia 11pt #4A5A70（居副位）
"""
from pathlib import Path

from docx import Document
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import parse_xml
from docx.oxml.ns import nsdecls, qn
from docx.shared import Cm, Pt, RGBColor

BASE = Path(__file__).resolve().parent
FIG = BASE / "figures"
# 发布产物统一落在 dist/。
# 早期版本输出到 whitepaper/ 根目录，于是同一份 DOCX 会在根目录与 dist/ 各存一份，
# 两个副本容易不同步。现在只保留 dist/ 这一个规范位置。
DIST = BASE / "dist"
DIST.mkdir(exist_ok=True)
OUT = DIST / "意义智能_元智能理论与分层应用_公开发布白皮书_v1.0_中英双语.docx"

# ---------------------------------------------------------------- 排版常量
NAVY = "1B3A5C"
BODY = "222222"
GRAY = "7A8CA3"
SECOND = "4A5A70"
LIGHT = "EEF2F7"

HEI = "黑体"
SONG = "宋体"
LATIN_H = "Arial"
# 正典用 Georgia，但 Georgia 默认旧体数字（old-style figures）：
# 「300」会被渲染成「3oo」、「V7.0」成「V7.o」、「ASD-STE100」成「ASD-STE1oo」。
# 本文档 E0–E6 / L1–L6 / 百分比密集，故改用同为衬线、但使用等高数字的 Cambria。
# 若要严格回到正典外观，把下面一行改回 "Georgia" 即可。
LATIN_B = "Cambria"

CONTENT_W = Cm(21 - 3.17 * 2)


def set_font(run, size, color, bold=False, east=SONG, latin=LATIN_B, italic=False):
    run.font.size = Pt(size)
    run.font.bold = bold
    run.font.italic = italic
    run.font.color.rgb = RGBColor.from_string(color)
    run.font.name = latin
    rPr = run._element.get_or_add_rPr()
    rf = rPr.get_or_add_rFonts()
    rf.set(qn("w:eastAsia"), east)
    rf.set(qn("w:ascii"), latin)
    rf.set(qn("w:hAnsi"), latin)
    return run


def add_rich(p, text, size, color, east=SONG, latin=LATIN_B, base_bold=False):
    """把 **粗体** 标记解析为真正的加粗 run。

    之前直接在正文里写 **层级更高**，DOCX 不认 Markdown，星号被原样渲染到纸面上。
    """
    for i, seg in enumerate(text.split("**")):
        if not seg:
            continue
        set_font(p.add_run(seg), size, color, base_bold or (i % 2 == 1), east, latin)
    return p


def para(doc, spacing=1.5, before=0, after=6, indent=None, align=None):
    p = doc.add_paragraph()
    pf = p.paragraph_format
    pf.line_spacing = spacing
    pf.space_before = Pt(before)
    pf.space_after = Pt(after)
    if indent is not None:
        pf.first_line_indent = indent
    if align is not None:
        p.alignment = align
    return p


def title(doc, zh, en, size=20, after=10, before=0, align=None):
    p = para(doc, spacing=1.15, before=before, after=after, align=align)
    set_font(p.add_run(zh), size, NAVY, True, HEI, LATIN_H)
    if en:
        p2 = para(doc, spacing=1.15, before=0, after=after + 4, align=align)
        set_font(p2.add_run(en), size * 0.62, GRAY, False, HEI, LATIN_H)
    return p


def h(doc, zh, en="", size=15, before=16, after=7):
    p = para(doc, spacing=1.25, before=before, after=after)
    set_font(p.add_run(zh), size, NAVY, True, HEI, LATIN_H)
    if en:
        set_font(p.add_run("\u3000" + en), size * 0.72, NAVY, False, HEI, LATIN_H)
    return p


def label(doc, zh, en=""):
    p = para(doc, spacing=1.2, before=10, after=3)
    add_rich(p, zh, 11, NAVY, HEI, LATIN_H, base_bold=True)
    if en:
        set_font(p.add_run("\u3000" + en), 9, GRAY, False, HEI, LATIN_H)
    return p


def body(doc, zh, en=None):
    p = para(doc, spacing=1.5, after=4 if en else 7, indent=Cm(0.74))
    add_rich(p, zh, 12, BODY, SONG, LATIN_B)
    if en:
        q = para(doc, spacing=1.35, after=8, indent=Cm(0.74))
        add_rich(q, en, 11, SECOND, SONG, LATIN_B)
    return p


def bullet(doc, zh, en=None):
    p = para(doc, spacing=1.4, after=2)
    p.paragraph_format.left_indent = Cm(0.74)
    add_rich(p, "· " + zh, 11.5, BODY, SONG, LATIN_B)
    if en:
        q = para(doc, spacing=1.3, after=6)
        q.paragraph_format.left_indent = Cm(1.10)
        add_rich(q, en, 10.5, SECOND, SONG, LATIN_B)
    return p


def note(doc, zh, en=None):
    # 必须走 add_rich：否则正文里的 **强调** 标记会被原样打印到纸面上。
    # （body/bullet 早已支持；note 曾漏掉，导致 L3 新增的光谱带注释把星号印了出来。）
    p = para(doc, spacing=1.3, before=4, after=3)
    add_rich(p, zh, 9.5, SECOND, SONG, LATIN_B)
    if en:
        q = para(doc, spacing=1.25, after=8)
        add_rich(q, en, 9, GRAY, SONG, LATIN_B)
    return p


def shade(cell, fill):
    cell._tc.get_or_add_tcPr().append(
        parse_xml(r'<w:shd {} w:val="clear" w:fill="{}"/>'.format(nsdecls("w"), fill)))


def table(doc, headers, rows, widths=None, caption=None, caption_en=None):
    if caption:
        p = para(doc, spacing=1.2, before=12, after=4)
        set_font(p.add_run(caption), 11, NAVY, True, HEI, LATIN_H)
        if caption_en:
            set_font(p.add_run("  |  " + caption_en), 8.5, GRAY, False, HEI, LATIN_H)
    t = doc.add_table(rows=1, cols=len(headers))
    t.style = "Table Grid"
    t.alignment = WD_TABLE_ALIGNMENT.CENTER
    # 固定列宽：不锁 layout 的话 Word/LO 会按内容自动重排，
    # 导致表头如「理论锚点（正典）」被挤成两行。
    t.autofit = False
    t._tbl.tblPr.append(
        parse_xml(r'<w:tblLayout {} w:type="fixed"/>'.format(nsdecls("w"))))
    hdr = t.rows[0].cells
    for i, htxt in enumerate(headers):
        shade(hdr[i], NAVY)
        cp = hdr[i].paragraphs[0]
        cp.paragraph_format.line_spacing = 1.15
        cp.paragraph_format.space_after = Pt(1)
        cp.paragraph_format.space_before = Pt(1)
        set_font(cp.add_run(htxt), 9.5, "FFFFFF", True, HEI, LATIN_H)
    for ri, row in enumerate(rows):
        cells = t.add_row().cells
        for i, val in enumerate(row):
            if ri % 2 == 1:
                shade(cells[i], LIGHT)
            cp = cells[i].paragraphs[0]
            cp.paragraph_format.line_spacing = 1.2
            cp.paragraph_format.space_after = Pt(1)
            cp.paragraph_format.space_before = Pt(1)
            # 同样走 add_rich：表格单元格里的 **强调** 也必须变成真加粗
            add_rich(cp, str(val), 9, BODY, SONG, LATIN_B)
    if widths:
        for i, w in enumerate(widths):
            t.columns[i].width = Cm(w)
        for r in t.rows:
            for i, w in enumerate(widths):
                r.cells[i].width = Cm(w)
    para(doc, spacing=1.0, after=8)
    return t


def figure(doc, path, width_cm, caption, caption_en):
    p = para(doc, spacing=1.0, before=6, after=3, align=WD_ALIGN_PARAGRAPH.CENTER)
    p.add_run().add_picture(str(path), width=Cm(width_cm))
    c = para(doc, spacing=1.25, after=10, align=WD_ALIGN_PARAGRAPH.CENTER)
    set_font(c.add_run(caption), 10, NAVY, True, HEI, LATIN_H)
    c2 = para(doc, spacing=1.2, after=12, align=WD_ALIGN_PARAGRAPH.CENTER)
    set_font(c2.add_run(caption_en), 8.5, GRAY, False, HEI, LATIN_H)
